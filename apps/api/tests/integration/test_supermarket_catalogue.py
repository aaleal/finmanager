"""Protects: the M1 rubric bullets *«every line resolves to a MasterProduct and
inherits its category»*, *«a brand-new product arrives with an AUTO suggestion a
human can promote to VALIDATED once, for every past and future line»*, and the
four category-governance rules (rename, reparent, merge, retire).
"""

from __future__ import annotations

import io
from decimal import Decimal
from typing import Any

import openpyxl
import pytest
from app.core.errors import Conflict, ValidationError
from app.models import Entity, Merchant, User
from app.models.core import Category
from app.models.products import MasterProduct, ProductAlias
from app.models.supermarket import SupermarketReceipt, SupermarketReceiptItem
from app.services.supermarket import catalogue, products_service
from app.services.supermarket import service as receipts
from sqlalchemy import select
from sqlalchemy.orm import Session

pytestmark = pytest.mark.integration


@pytest.fixture
def tree(db: Session) -> dict[str, Category]:
    l1 = products_service.create_category(db, display_name_pt="Laticínios", parent_id=None)
    l2 = products_service.create_category(db, display_name_pt="Queijo Fresco", parent_id=l1.id)
    l3 = products_service.create_category(db, display_name_pt="Queijo Mascarpone", parent_id=l2.id)
    other = products_service.create_category(db, display_name_pt="Mercearia", parent_id=None)
    return {"l1": l1, "l2": l2, "l3": l3, "other": other}


@pytest.fixture
def merchant(db: Session) -> Merchant:
    row = Merchant(name="Pingo Doce", kind="RETAIL", aliases=[])
    db.add(row)
    db.flush()
    return row


def make_product(db: Session, tree: dict[str, Category], name: str = "Mascarpone Galbani"):  # type: ignore[no-untyped-def]
    return products_service.create_product(
        db, canonical_name=name, category_id=tree["l3"].id, category_status="AUTO"
    )


# --- Ancestry ------------------------------------------------------------------


def test_assigning_the_deepest_category_maintains_its_ancestors(
    db: Session, tree: dict[str, Category]
) -> None:
    product = make_product(db, tree)
    assert product.category_id == tree["l3"].id
    assert product.category_l1_id == tree["l1"].id
    assert product.category_l2_id == tree["l2"].id
    assert product.category_l3_id == tree["l3"].id


def test_an_l2_only_assignment_is_a_real_answer_not_a_gap(
    db: Session, tree: dict[str, Category]
) -> None:
    """The household's own sheet leaves L3 blank on nearly half its rows."""
    product = products_service.create_product(
        db, canonical_name="Queijo qualquer", category_id=tree["l2"].id
    )
    assert product.category_l2_id == tree["l2"].id
    assert product.category_l3_id is None


def test_the_full_path_disambiguates_a_leaf_that_lives_under_two_parents(
    db: Session, tree: dict[str, Category]
) -> None:
    assert (
        products_service.category_path(db, tree["l3"].id)
        == "Laticínios › Queijo Fresco › Queijo Mascarpone"
    )


# --- Category governance -------------------------------------------------------


def test_a_rename_leaves_every_referencing_row_untouched(
    db: Session, tree: dict[str, Category], owner: User
) -> None:
    product = make_product(db, tree)
    before = (product.category_id, product.category_l1_id, product.category_l3_id)

    products_service.rename_category(db, tree["l2"], "Queijos Frescos", actor_user_id=owner.id)

    assert (product.category_id, product.category_l1_id, product.category_l3_id) == before
    assert products_service.category_path(db, tree["l3"].id) == (
        "Laticínios › Queijos Frescos › Queijo Mascarpone"
    )


def test_a_reparent_recomputes_maintained_ancestors_on_every_affected_product(
    db: Session, tree: dict[str, Category], owner: User
) -> None:
    product = make_product(db, tree)
    affected = products_service.reparent_category(
        db, tree["l2"], tree["other"].id, actor_user_id=owner.id
    )

    assert affected == 1
    assert product.category_l1_id == tree["other"].id
    assert product.category_l2_id == tree["l2"].id
    assert product.category_l3_id == tree["l3"].id


def test_a_category_cannot_be_moved_inside_itself(
    db: Session, tree: dict[str, Category], owner: User
) -> None:
    with pytest.raises(ValidationError):
        products_service.reparent_category(db, tree["l1"], tree["l3"].id, actor_user_id=owner.id)


def test_retire_is_refused_while_the_category_is_in_use(
    db: Session, tree: dict[str, Category], owner: User
) -> None:
    make_product(db, tree)
    usage = products_service.impact(db, tree["l3"].id)
    assert usage.in_use is True
    assert usage.master_products == 1

    with pytest.raises(Conflict):
        products_service.retire_category(db, tree["l3"], actor_user_id=owner.id)


def test_merge_reassigns_every_row_then_retires_the_source(
    db: Session, tree: dict[str, Category], owner: User
) -> None:
    product = make_product(db, tree)
    target = products_service.create_category(
        db, display_name_pt="Queijo Creme", parent_id=tree["l2"].id
    )

    moved = products_service.merge_categories(
        db, source=tree["l3"], target=target, actor_user_id=owner.id
    )

    assert moved == 1
    assert product.category_id == target.id
    assert db.get(Category, tree["l3"].id).is_deleted is True

    # …and the freed category can now be retired, which is the way out of «in use».
    products_service.retire_category(db, tree["l3"], actor_user_id=owner.id)


def test_an_unused_category_retires_cleanly(
    db: Session, tree: dict[str, Category], owner: User
) -> None:
    products_service.retire_category(db, tree["other"], actor_user_id=owner.id)
    assert db.get(Category, tree["other"].id).is_deleted is True


# --- Whole-tree export / import (Decision #56) ----------------------------------


def _category_import_workbook(rows: list[list[Any]]) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Nível 1", "Nível 2", "Nível 3", "Eixo de marca"])
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_export_writes_one_row_per_node_with_its_full_path(
    db: Session, tree: dict[str, Category]
) -> None:
    payload = products_service.export_categories_workbook(db)
    workbook = openpyxl.load_workbook(io.BytesIO(payload))
    rows = list(workbook["Categorias"].iter_rows(values_only=True))[1:]

    assert ("Laticínios", None, None, None) in rows
    assert ("Laticínios", "Queijo Fresco", None, None) in rows
    assert ("Laticínios", "Queijo Fresco", "Queijo Mascarpone", None) in rows
    assert ("Mercearia", None, None, None) in rows


def test_import_only_creates_what_is_missing_and_is_safe_to_rerun(
    db: Session, tree: dict[str, Category], owner: User
) -> None:
    data = _category_import_workbook(
        [
            ["Laticínios", "Queijo Fresco", "Queijo Mascarpone", None],  # already there
            ["Laticínios", "Queijo Fresco", "Queijo Creme", None],  # new leaf, known parents
            ["Bebidas", "Águas", None, "Sim"],  # brand new branch
        ]
    )

    report = products_service.import_categories_workbook(db, data=data, actor_user_id=owner.id)
    assert report.errors == []
    assert report.existing == 1
    assert report.created == 3  # Queijo Creme + Bebidas + Águas

    aguas = db.scalar(select(Category).where(Category.display_name_pt == "Águas"))
    assert aguas is not None
    assert aguas.brand_axis is True
    assert aguas.level == 2

    # Re-running the exact same sheet creates nothing new.
    again = products_service.import_categories_workbook(db, data=data, actor_user_id=owner.id)
    assert again.created == 0
    assert again.existing == 3


def test_import_rejects_a_row_with_a_gap_before_a_filled_level(
    db: Session, tree: dict[str, Category], owner: User
) -> None:
    data = _category_import_workbook([[None, "Órfã", None, None]])

    report = products_service.import_categories_workbook(db, data=data, actor_user_id=owner.id)

    assert report.created == 0
    assert report.existing == 0
    assert [(e.row_number, e.message) for e in report.errors] == [
        (2, "Falta um nível antes de um nível preenchido.")
    ]


# --- Resolution and the learning loop ------------------------------------------


def test_a_confirmed_correction_sticks_for_the_same_merchant_text(
    db: Session, tree: dict[str, Category], merchant: Merchant
) -> None:
    product = make_product(db, tree)
    before = catalogue.resolve_description(
        db, merchant_id=merchant.id, description_norm="MASCARP GALB 250G"
    )
    assert before.master_product_id is None

    catalogue.learn(
        db,
        master_product_id=product.id,
        merchant_id=merchant.id,
        merchant_description="MASCARP GALB 250G",
    )
    after = catalogue.resolve_description(
        db, merchant_id=merchant.id, description_norm="MASCARP GALB 250G"
    )
    assert after.master_product_id == product.id


def test_alias_confidence_rises_with_every_confirmation(
    db: Session, tree: dict[str, Category], merchant: Merchant
) -> None:
    """Learn on the *first* correction — waiting for five discards the signal."""
    product = make_product(db, tree)
    first = catalogue.learn(
        db, master_product_id=product.id, merchant_id=merchant.id, merchant_description="X"
    )
    assert first.confidence == Decimal("0.850")

    second = catalogue.learn(
        db, master_product_id=product.id, merchant_id=merchant.id, merchant_description="X"
    )
    assert second.confidence == Decimal("0.950")
    assert second.correction_count == 2


def test_repointing_an_alias_restarts_the_trust_it_had_earned(
    db: Session, tree: dict[str, Category], merchant: Merchant
) -> None:
    first = make_product(db, tree, name="Mascarpone Galbani")
    second = make_product(db, tree, name="Mascarpone Continente")
    catalogue.learn(
        db, master_product_id=first.id, merchant_id=merchant.id, merchant_description="MASC"
    )
    catalogue.learn(
        db, master_product_id=first.id, merchant_id=merchant.id, merchant_description="MASC"
    )

    corrected = catalogue.learn(
        db, master_product_id=second.id, merchant_id=merchant.id, merchant_description="MASC"
    )
    assert corrected.master_product_id == second.id
    assert corrected.correction_count == 1


def test_an_empty_catalogue_resolves_nothing_and_says_why(db: Session, merchant: Merchant) -> None:
    match = catalogue.resolve_description(
        db, merchant_id=merchant.id, description_norm="QUALQUER COISA"
    )
    assert match.master_product_id is None
    assert match.reasons[0]["rule"] == "empty_catalogue"


# --- Confirming a category once, for every line --------------------------------


def test_confirming_a_category_settles_every_past_and_future_line(
    db: Session, entity: Entity, owner: User, tree: dict[str, Category], merchant: Merchant
) -> None:
    product = make_product(db, tree)
    receipt = SupermarketReceipt(
        entity_id=entity.id, merchant_id=merchant.id, status="NEEDS_REVIEW"
    )
    db.add(receipt)
    db.flush()
    for index in range(2):
        db.add(
            SupermarketReceiptItem(
                receipt_id=receipt.id,
                entity_id=entity.id,
                line_no=index + 1,
                description_raw="MASCARP GALB 250G",
                description_norm="MASCARP GALB 250G",
                master_product_id=product.id,
                unit_price_pvp_eur=Decimal("4.19"),
                paid_price_eur=Decimal("4.19"),
            )
        )
    db.flush()
    db.refresh(receipt)

    assert receipts.confirm_categories(db, receipt, actor_user_id=owner.id) == 1
    assert product.category_status == "VALIDATED"
    # The backlog is «products never checked», a list that shrinks.
    assert products_service.uncategorized_count(db) == 0


def test_reassigning_a_line_moves_the_category_and_teaches_the_alias(
    db: Session, entity: Entity, owner: User, tree: dict[str, Category], merchant: Merchant
) -> None:
    wrong = make_product(db, tree, name="Produto Errado")
    right = make_product(db, tree, name="Produto Certo")
    receipt = SupermarketReceipt(
        entity_id=entity.id, merchant_id=merchant.id, status="NEEDS_REVIEW"
    )
    db.add(receipt)
    db.flush()
    item = SupermarketReceiptItem(
        receipt_id=receipt.id,
        entity_id=entity.id,
        line_no=1,
        description_raw="TEXTO DO TALAO",
        description_norm="TEXTO DO TALAO",
        master_product_id=wrong.id,
        unit_price_pvp_eur=Decimal("1.00"),
        paid_price_eur=Decimal("1.00"),
    )
    db.add(item)
    db.flush()

    receipts.reassign_item_product(db, item, master_product_id=right.id, actor_user_id=owner.id)

    assert item.master_product_id == right.id
    alias = db.scalar(
        select(ProductAlias).where(
            ProductAlias.merchant_id == merchant.id,
            ProductAlias.description_norm == "TEXTO DO TALAO",
        )
    )
    assert alias is not None and alias.master_product_id == right.id


def test_merging_two_products_moves_their_lines_and_aliases(
    db: Session, entity: Entity, owner: User, tree: dict[str, Category], merchant: Merchant
) -> None:
    source = make_product(db, tree, name="Mascarpone Galbani 250")
    target = make_product(db, tree, name="Mascarpone Galbani")
    catalogue.learn(
        db, master_product_id=source.id, merchant_id=merchant.id, merchant_description="MASC 250"
    )
    receipt = SupermarketReceipt(
        entity_id=entity.id, merchant_id=merchant.id, status="NEEDS_REVIEW"
    )
    db.add(receipt)
    db.flush()
    item = SupermarketReceiptItem(
        receipt_id=receipt.id,
        entity_id=entity.id,
        line_no=1,
        description_raw="MASC 250",
        description_norm="MASC 250",
        master_product_id=source.id,
        unit_price_pvp_eur=Decimal("4.19"),
        paid_price_eur=Decimal("4.19"),
    )
    db.add(item)
    db.flush()

    products_service.merge_products(db, source=source, target=target, actor_user_id=owner.id)
    db.flush()

    assert db.get(MasterProduct, source.id).is_deleted is True
    assert db.get(SupermarketReceiptItem, item.id).master_product_id == target.id
    alias = db.scalar(select(ProductAlias).where(ProductAlias.description_norm == "MASC 250"))
    assert alias is not None and alias.master_product_id == target.id
