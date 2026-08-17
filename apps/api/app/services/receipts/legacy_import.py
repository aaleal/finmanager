"""FR-1.19 - the ``SUPERMARKET_YYYY`` migration.

**The sheet is validated, not trusted.** It was maintained by hand over years and
is wrong in places, so every imported receipt is scored like a parsed one: it
gets a confidence and decision reasons, and a group that does not reconcile lands
in ``NEEDS_REVIEW`` instead of being written as fact.

Two things about the source are load-bearing and were measured against the real
2025 file (2,429 rows), not assumed:

**Undo the Fs timestamp nudge first, then group by the timestamp itself.** Fs rows
were recorded about a second after the invoice they belong to - an artefact of
how the sheet was filled in, not a fact about the purchase. Measured: 401 of 425
Fs rows sit at exactly +1 s from the nearest earlier non-Fs row of the same shop
and day, and 17 sit at +0 s. Snapping beats a blanket -1 s precisely because of
those 17.

**Never group by ``Preço Fatura``.** It is a ``SUMIF`` over ``Full_Date``, so on a
nudged Fs row it sums only itself - self-referential and meaningless. Grouping on
it invents a separate one-line invoice for the Fs article, the exact opposite of
the association the Fs model exists to preserve.

Measured outcome of the algorithm below on the real file: 288 receipts, 6 groups
that are only Fs, 280 reconciling within €0.02, 1,564 distinct products.
"""

from __future__ import annotations

import datetime as dt
import io
import uuid
from collections import defaultdict
from dataclasses import dataclass, field
from decimal import Decimal
from typing import Any

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session as DbSession

from app.core.errors import ValidationError
from app.core.money import ZERO, to_eur
from app.models.core import ImportBatch, Merchant
from app.models.products import MasterProduct
from app.models.receipts import Receipt, ReceiptItem
from app.services.receipts import arithmetic, catalogue, prices_service, products_service
from app.services.receipts.normalize import normalize_description, normalize_merchant_name

#: Fs rows were nudged by about a second; anything further apart is a real gap.
SNAP_WINDOW = dt.timedelta(seconds=2)
TOLERANCE = Decimal("0.02")

#: Outside this module's scope for now. Excluded and **counted**, so the decision
#: stays visible and reversible - widening the scope later is a re-run.
NON_GROCERY = {
    "IKEA",
    "LEROY",
    "WELLS",
    "NORMAL",
    "ACTION",
    "EL CORTE INGLES",
    "EL CORT INGLES",
}

REQUIRED_COLUMNS = ("Full_Date", "Supermercado", "Descrição", "Price", "Price_Final")


@dataclass(slots=True)
class ImportReport:
    import_batch_id: uuid.UUID
    row_count: int = 0
    skipped_non_grocery_rows: int = 0
    skipped_non_grocery_merchants: list[str] = field(default_factory=list)
    receipts_created: int = 0
    receipts_reconciled: int = 0
    receipts_needing_review: int = 0
    fs_rows: int = 0
    fs_rows_snapped: int = 0
    all_fs_groups: int = 0
    products_created: int = 0
    aliases_created: int = 0
    merchants_created: int = 0
    exceptions: list[dict[str, Any]] = field(default_factory=list)

    def as_dict(self) -> dict[str, Any]:
        return {
            "import_batch_id": str(self.import_batch_id),
            "row_count": self.row_count,
            "skipped_non_grocery_rows": self.skipped_non_grocery_rows,
            "skipped_non_grocery_merchants": sorted(set(self.skipped_non_grocery_merchants)),
            "receipts_created": self.receipts_created,
            "receipts_reconciled": self.receipts_reconciled,
            "receipts_needing_review": self.receipts_needing_review,
            "fs_rows": self.fs_rows,
            "fs_rows_snapped": self.fs_rows_snapped,
            "all_fs_groups": self.all_fs_groups,
            "products_created": self.products_created,
            "aliases_created": self.aliases_created,
            "merchants_created": self.merchants_created,
            "exceptions": self.exceptions[:200],
        }


# --- Reading ------------------------------------------------------------------


def _decimal(value: Any) -> Decimal:
    if value is None or value == "":
        return ZERO
    return to_eur(Decimal(str(value))) or ZERO


def _is_fs(row: dict[str, Any]) -> bool:
    """``F`` or ``f`` - matched case-insensitively; 425 of the 2,429 rows carry one."""
    return str(row.get("Flags") or "").strip().upper() == "F"


def _weight(value: Any) -> Decimal | None:
    if value in (None, ""):
        return None
    weight = Decimal(str(value))
    return weight if weight > 0 else None


def read_rows(data: bytes) -> list[dict[str, Any]]:
    workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    iterator = sheet.iter_rows(values_only=True)
    try:
        header = [str(cell) if cell is not None else "" for cell in next(iterator)]
    except StopIteration as exc:
        raise ValidationError("A folha está vazia.") from exc

    missing = [column for column in REQUIRED_COLUMNS if column not in header]
    if missing:
        raise ValidationError(f"Colunas em falta na folha: {', '.join(missing)}.")

    rows: list[dict[str, Any]] = []
    for raw in iterator:
        row = dict(zip(header, raw, strict=False))
        if row.get("Full_Date") is None or row.get("Supermercado") is None:
            continue
        rows.append(row)
    workbook.close()
    return rows


def snap_fs_timestamps(rows: list[dict[str, Any]]) -> int:
    """Put every Fs row back on the invoice it belongs to, then leave it alone."""
    anchors: dict[tuple[str, dt.date], list[dt.datetime]] = defaultdict(list)
    for row in rows:
        if not _is_fs(row):
            anchor: dt.datetime = row["Full_Date"]
            anchors[(normalize_merchant_name(str(row["Supermercado"])), anchor.date())].append(
                anchor
            )

    snapped = 0
    for row in rows:
        stamp: dt.datetime = row["Full_Date"]
        if _is_fs(row):
            key = (normalize_merchant_name(str(row["Supermercado"])), stamp.date())
            nearest = min(
                (candidate for candidate in anchors.get(key, [])),
                key=lambda candidate: abs(candidate - stamp),
                default=None,
            )
            if nearest is not None and abs(nearest - stamp) <= SNAP_WINDOW:
                stamp = nearest
                snapped += 1
        row["_timestamp"] = stamp
    return snapped


def group_rows(rows: list[dict[str, Any]]) -> dict[tuple[str, dt.datetime], list[dict[str, Any]]]:
    groups: dict[tuple[str, dt.datetime], list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        groups[(normalize_merchant_name(str(row["Supermercado"])), row["_timestamp"])].append(row)
    return groups


# --- Writing ------------------------------------------------------------------


def _find_or_create_merchant(
    db: DbSession, name: str, cache: dict[str, Merchant], report: ImportReport
) -> Merchant:
    """``Mercadona`` and ``mercadona`` are one shop; so are stray-whitespace twins."""
    key = normalize_merchant_name(name)
    if key in cache:
        return cache[key]

    for merchant in db.scalars(select(Merchant).where(Merchant.is_deleted.is_(False))).all():
        names = [merchant.name, *(str(a) for a in (merchant.aliases or []))]
        if any(normalize_merchant_name(candidate) == key for candidate in names):
            cache[key] = merchant
            return merchant

    merchant = Merchant(name=name.strip(), kind="RETAIL", aliases=[])
    db.add(merchant)
    db.flush()
    report.merchants_created += 1
    cache[key] = merchant
    return merchant


def _find_or_create_product(
    db: DbSession, row: dict[str, Any], cache: dict[str, MasterProduct], report: ImportReport
) -> MasterProduct:
    description = str(row["Descrição"] or "").strip()
    key = normalize_description(description)
    if key in cache:
        return cache[key]

    existing = db.scalar(
        select(MasterProduct).where(
            MasterProduct.canonical_name == description, MasterProduct.is_deleted.is_(False)
        )
    )
    if existing is not None:
        cache[key] = existing
        return existing

    # The sheet's own categories are per-row and demonstrably unreliable - one
    # fixture row files a cider under «Talho › Vaca › Almondegas». They import as
    # an AUTO *suggestion* on the product, never as VALIDATED.
    category_id = _category_from_sheet(db, row)
    product = products_service.create_product(
        db,
        canonical_name=description[:250],
        category_id=category_id,
        category_status="AUTO",
        category_confidence=Decimal("0.400") if category_id else None,
        sold_by_weight=False,
    )
    report.products_created += 1
    cache[key] = product
    return product


def _category_from_sheet(db: DbSession, row: dict[str, Any]) -> uuid.UUID | None:
    from app.models.core import Category

    parent_id: uuid.UUID | None = None
    deepest: uuid.UUID | None = None
    for column, level in (("Categoria", 1), ("Categoria_2", 2), ("Categoria_3", 3)):
        name = str(row.get(column) or "").strip()
        if not name:
            break
        node = db.scalar(
            select(Category).where(
                Category.domain == "GROCERY",
                Category.level == level,
                Category.parent_id.is_(parent_id)
                if parent_id is None
                else Category.parent_id == parent_id,
                Category.display_name_pt.ilike(name),
                Category.is_deleted.is_(False),
            )
        )
        if node is None:
            break
        deepest, parent_id = node.id, node.id
    return deepest


def import_workbook(
    db: DbSession,
    *,
    data: bytes,
    entity_id: uuid.UUID,
    actor_user_id: uuid.UUID | None,
    document_id: uuid.UUID | None = None,
    include_non_grocery: bool = False,
) -> ImportReport:
    """Idempotent: a group already imported for this entity is updated, not doubled."""
    rows = read_rows(data)
    batch = ImportBatch(
        entity_id=entity_id,
        module="receipts",
        source_type="MANUAL",
        file_document_id=document_id,
        status="PROCESSING",
        row_count=len(rows),
    )
    db.add(batch)
    db.flush()

    report = ImportReport(import_batch_id=batch.id, row_count=len(rows))
    report.fs_rows = sum(1 for row in rows if _is_fs(row))
    report.fs_rows_snapped = snap_fs_timestamps(rows)

    kept: list[dict[str, Any]] = []
    for row in rows:
        name = normalize_merchant_name(str(row["Supermercado"]))
        if not include_non_grocery and name in NON_GROCERY:
            report.skipped_non_grocery_rows += 1
            report.skipped_non_grocery_merchants.append(str(row["Supermercado"]).strip())
            continue
        kept.append(row)

    merchant_cache: dict[str, Merchant] = {}
    product_cache: dict[str, MasterProduct] = {}

    for (_, timestamp), group in sorted(group_rows(kept).items(), key=lambda pair: pair[0][1]):
        printed = [row for row in group if not _is_fs(row)]
        if not printed:
            # Fs articles with no invoice to attach to. Imported as receipts with
            # total 0,00 and flagged, rather than inventing a parent for them.
            report.all_fs_groups += 1
        _import_group(
            db,
            group,
            timestamp=timestamp,
            entity_id=entity_id,
            batch=batch,
            merchant_cache=merchant_cache,
            product_cache=product_cache,
            report=report,
            actor_user_id=actor_user_id,
        )

    batch.status = "COMPLETED"
    batch.success_count = report.receipts_created
    batch.error_count = len(report.exceptions)
    batch.completed_at = dt.datetime.now(dt.UTC)
    db.flush()
    return report


def _import_group(
    db: DbSession,
    group: list[dict[str, Any]],
    *,
    timestamp: dt.datetime,
    entity_id: uuid.UUID,
    batch: ImportBatch,
    merchant_cache: dict[str, Merchant],
    product_cache: dict[str, MasterProduct],
    report: ImportReport,
    actor_user_id: uuid.UUID | None,
) -> None:
    merchant = _find_or_create_merchant(db, str(group[0]["Supermercado"]), merchant_cache, report)
    printed = [row for row in group if not _is_fs(row)]
    fs_rows = [row for row in group if _is_fs(row)]

    # An Fs article with no value is meaningless, and the sheet has a few. They
    # are reported as exceptions rather than silently repaired or silently dropped.
    valued_fs: list[dict[str, Any]] = []
    unvalued_fs: list[dict[str, Any]] = []
    for row in fs_rows:
        (valued_fs if _decimal(row.get("Price")) > ZERO else unvalued_fs).append(row)
    for row in unvalued_fs:
        report.exceptions.append(
            {
                "rule": "fs_row_without_value",
                "merchant": str(group[0]["Supermercado"]).strip(),
                "timestamp": timestamp.isoformat(),
                "description": str(row.get("Descrição") or "").strip(),
            }
        )
    fs_rows = valued_fs

    # Taken from the non-Fs rows only: on a nudged Fs row the sheet's own
    # «Preço Fatura» summed nothing but itself.
    claimed_total = max((_decimal(row.get("Preço Fatura")) for row in printed), default=ZERO)

    receipt = db.scalar(
        select(Receipt).where(
            Receipt.entity_id == entity_id,
            Receipt.merchant_id == merchant.id,
            Receipt.purchased_at == timestamp,
            Receipt.is_deleted.is_(False),
        )
    )
    created = receipt is None
    if receipt is None:
        receipt = Receipt(
            entity_id=entity_id,
            merchant_id=merchant.id,
            import_batch_id=batch.id,
            status="UPLOADED",
        )
        db.add(receipt)
        db.flush()
    else:
        receipt.items.clear()
        db.flush()

    receipt.purchased_at = timestamp
    receipt.purchase_date = timestamp.date()
    receipt.total_eur = claimed_total
    receipt.item_count = len(printed)

    gross = sum((_decimal(row.get("Price")) for row in printed), ZERO)
    promos = sum((_decimal(row.get("PromoInd")) for row in printed), ZERO)
    # «PromoGlob» was filled on 8 % of rows, so it is recomputed deterministically
    # for every row from the receipt-level residual rather than trusted.
    invoice_credit = max(gross - promos - claimed_total, ZERO)
    receipt.total_discount_eur = invoice_credit

    allocations = arithmetic.prorate_invoice_discount(
        [
            arithmetic.ProrationInput(_decimal(row.get("Price")), _decimal(row.get("PromoInd")))
            for row in printed
        ],
        invoice_credit,
    )

    items: list[ReceiptItem] = []
    for index, (row, allocated) in enumerate(zip(printed, allocations, strict=True), start=1):
        items.append(
            _row_to_item(
                db,
                row,
                receipt=receipt,
                merchant=merchant,
                line_no=index,
                allocated=allocated,
                product_cache=product_cache,
                report=report,
            )
        )
    for row in fs_rows:
        items.append(
            _row_to_item(
                db,
                row,
                receipt=receipt,
                merchant=merchant,
                line_no=None,
                allocated=ZERO,
                product_cache=product_cache,
                report=report,
                is_fs=True,
            )
        )
    receipt.items = items
    db.flush()

    totals = arithmetic.receipt_totals(
        [
            arithmetic.ItemTotalsInput(
                paid_price_eur=item.paid_price_eur,
                unit_price_pvp_eur=item.unit_price_pvp_eur,
                promo_discount_eur=item.promo_discount_eur,
                quantity=item.quantity,
                is_fs=item.is_fs,
            )
            for item in items
        ],
        total_eur=claimed_total,
        total_discount_eur=invoice_credit,
        tolerance_eur=TOLERANCE,
    )

    reasons = [
        {
            "rule": "legacy_import",
            "detail": f"Importado da folha SUPERMARKET, {len(printed)} artigos impressos.",
            "score": None,
        }
    ]
    if totals.is_reconciled:
        receipt.status = "AUTO_ACCEPTED"
        receipt.confidence = Decimal("0.900")
        report.receipts_reconciled += 1
    else:
        receipt.status = "NEEDS_REVIEW"
        receipt.confidence = Decimal("0.450")
        reasons.append(
            {
                "rule": "arithmetic_mismatch",
                "detail": (
                    f"Σ linhas {totals.computed_total_eur} vs total da folha {claimed_total}."
                ),
                "score": "-0.500",
            }
        )
        report.receipts_needing_review += 1
        report.exceptions.append(
            {
                "merchant": merchant.name,
                "timestamp": timestamp.isoformat(),
                "computed_total_eur": str(totals.computed_total_eur),
                "claimed_total_eur": str(claimed_total),
                "item_count": len(printed),
            }
        )
    if not printed:
        reasons.append(
            {
                "rule": "orphan_fs_group",
                "detail": "Artigos Fs sem fatura à qual pertencer.",
                "score": None,
            }
        )
    receipt.decision_reasons = reasons

    # Nine years of €/kg history is the whole point of importing the sheet, so the
    # observations are frozen here rather than waiting for someone to re-confirm.
    prices_service.record_observations(db, receipt)

    if created:
        report.receipts_created += 1


def _row_to_item(
    db: DbSession,
    row: dict[str, Any],
    *,
    receipt: Receipt,
    merchant: Merchant,
    line_no: int | None,
    allocated: Decimal,
    product_cache: dict[str, MasterProduct],
    report: ImportReport,
    is_fs: bool = False,
) -> ReceiptItem:
    description = str(row["Descrição"] or "").strip()
    product = _find_or_create_product(db, row, product_cache, report)

    alias_before = product_cache.get(f"alias:{merchant.id}:{normalize_description(description)}")
    if alias_before is None:
        catalogue.learn(
            db,
            master_product_id=product.id,
            merchant_id=merchant.id,
            merchant_description=description,
        )
        product_cache[f"alias:{merchant.id}:{normalize_description(description)}"] = product
        report.aliases_created += 1

    pvp = _decimal(row.get("Price"))
    promo = _decimal(row.get("PromoInd"))
    # A blank and a literal zero mean the same thing here: no weight was recorded.
    # Storing 0 would suppress €/kg with «peso inválido» instead of «sem peso».
    weight = _weight(row.get("Peso"))
    listed = _weight(row.get("Peso (proposta)"))

    return ReceiptItem(
        receipt_id=receipt.id,
        entity_id=receipt.entity_id,
        line_no=line_no,
        description_raw=description[:300],
        description_norm=normalize_description(description)[:300],
        master_product_id=product.id,
        quantity=Decimal("1"),
        unit="UN",
        quantity_canonical=Decimal("1"),
        unit_canonical="UN",
        weight_observed_kg=weight,
        weight_listed_kg=listed,
        unit_price_pvp_eur=pvp,
        promo_discount_eur=ZERO if is_fs else promo,
        invoice_allocated_discount_eur=ZERO if is_fs else allocated,
        paid_price_eur=ZERO
        if is_fs
        else arithmetic.paid_from_components(
            unit_price_pvp_eur=pvp,
            promo_discount_eur=promo,
            invoice_allocated_discount_eur=allocated,
        ),
        is_fs=is_fs,
        notional_value_source="MANUAL" if is_fs else None,
        # «Notas» and «Status» both hold sparse free text, not an enum.
        notes=" · ".join(
            part
            for part in (str(row.get("Notas") or "").strip(), str(row.get("Status") or "").strip())
            if part
        )
        or None,
        legacy_row_ref=int(row["ID"]) if str(row.get("ID") or "").strip().isdigit() else None,
        confidence=Decimal("0.700"),
        decision_reasons=[
            {"rule": "legacy_import", "detail": "Linha importada da folha.", "score": "0.700"}
        ],
    )
