"""Workbook export for Module 9 (M9.1 FR-9.14).

One file, four sheets — copies, sets, storage locations and an auxiliary
reference sheet — so the collection can leave the application without an API
client. Read-only: it queries through the same service layer and writes
nothing.
"""

from __future__ import annotations

import datetime as dt
import io
import uuid
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session as DbSession
from sqlalchemy.orm import selectinload

from app.core.money import ZERO, appreciation_eur, roi_pct
from app.models.household import Entity
from app.models.lego import LegoSetInstance, LegoSetModel
from app.services import lego_service, settings_service

BUILD_STATE_PT = {"BUILT": "Montado", "DISASSEMBLED": "Desmontado"}
CONDITION_PT = {
    "SEALED": "Selado",
    "NEW": "Novo",
    "GOOD": "Bom",
    "WORN": "Usado",
    "DAMAGED": "Danificado",
}
SOURCE_PT = {
    "CONTINENTE": "Continente",
    "AMAZON": "Amazon",
    "OTHER_STORE": "Outra loja",
    "SECONDHAND": "Em segunda mão",
    "GIFT": "Prenda",
    "OTHER": "Outro",
}
OWNERSHIP_PT = {"IN_COLLECTION": "Na coleção", "SOLD": "Vendido", "GIFTED": "Oferecido"}

MONEY_FORMAT = '#,##0.00 "€"'
PERCENT_FORMAT = '0.00 "%"'
DATE_FORMAT = "dd/mm/yyyy"


def _write_sheet(
    workbook: Workbook,
    title: str,
    headers: list[str],
    rows: list[list[Any]],
    formats: dict[int, str],
) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        sheet.append(row)

    for index, format_code in formats.items():
        for cell in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
            cell[0].number_format = format_code

    for index, header in enumerate(headers, start=1):
        widest = max((len(str(row[index - 1] or "")) for row in rows), default=0)
        sheet.column_dimensions[get_column_letter(index)].width = min(
            max(len(header), widest) + 2, 45
        )

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def build_workbook(
    db: DbSession, *, entity_ids: list[uuid.UUID], active_entity_id: uuid.UUID | None
) -> bytes:
    entity_names = _entity_names(db, entity_ids)
    scope_ids = [active_entity_id] if active_entity_id else entity_ids

    workbook = Workbook()
    default_sheet = workbook.active
    if default_sheet is not None:
        workbook.remove(default_sheet)

    _copies_sheet(db, workbook, scope_ids, entity_names)
    _sets_sheet(db, workbook, scope_ids, entity_names)
    _locations_sheet(db, workbook)
    _helpers_sheet(workbook)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _entity_names(db: DbSession, entity_ids: list[uuid.UUID]) -> dict[uuid.UUID, str]:
    rows = db.execute(select(Entity.id, Entity.name).where(Entity.id.in_(entity_ids))).all()
    return {row[0]: row[1] for row in rows}


def _copies_sheet(
    db: DbSession,
    workbook: Workbook,
    scope_ids: list[uuid.UUID],
    entity_names: dict[uuid.UUID, str],
) -> None:
    copies = list(
        db.scalars(
            select(LegoSetInstance)
            .options(
                selectinload(LegoSetInstance.model),
                selectinload(LegoSetInstance.storage_location),
            )
            .join(LegoSetModel)
            .where(
                LegoSetInstance.entity_id.in_(scope_ids),
                LegoSetInstance.is_deleted.is_(False),
                LegoSetModel.is_deleted.is_(False),
            )
            .order_by(LegoSetModel.name)
        )
        .unique()
        .all()
    )

    rows: list[list[Any]] = []
    for copy in copies:
        model = copy.model
        location = copy.storage_location
        rows.append(
            [
                model.set_number or ("MOC" if model.is_custom else ""),
                model.name,
                model.theme,
                model.subtheme,
                model.release_date,
                model.retirement_date,
                model.piece_count,
                model.minifig_count,
                entity_names.get(copy.entity_id, ""),
                location.area if location else "",
                location.container if location else "",
                BUILD_STATE_PT.get(copy.build_state or "", ""),
                CONDITION_PT.get(copy.condition or "", ""),
                "Sim" if copy.has_box else "Não",
                "Sim" if copy.has_instructions else "Não",
                copy.missing_parts or "",
                "Não" if copy.missing_parts and copy.missing_parts.strip() else "Sim",
                OWNERSHIP_PT.get(copy.ownership_status, copy.ownership_status),
                copy.acquisition_date,
                SOURCE_PT.get(copy.acquisition_source or "", ""),
                copy.acquisition_cost_eur,
                model.rrp_eur,
                model.current_value_eur,
                appreciation_eur(copy.acquisition_cost_eur, model.current_value_eur),
                roi_pct(copy.acquisition_cost_eur, model.current_value_eur),
                appreciation_eur(model.rrp_eur, model.current_value_eur),
                roi_pct(model.rrp_eur, model.current_value_eur),
                model.value_updated_at,
                copy.sale_price_eur,
                copy.sale_date,
                copy.notes or "",
                "Sim" if copy.is_fs else "Não",
                "Sim" if copy.is_potential_gift else "Não",
            ]
        )

    _write_sheet(
        workbook,
        "Cópias",
        [
            "Número",
            "Nome",
            "Tema",
            "Subtema",
            "Lançamento",
            "Retirado em",
            "Peças",
            "Minifiguras",
            "Entidade",
            "Área",
            "Contentor",
            "Estado de construção",
            "Condição",
            "Tem caixa",
            "Tem instruções",
            "Peças em falta",
            "Completo",
            "Propriedade",
            "Data de aquisição",
            "Origem",
            "Custo (€)",
            "PVP original (€)",
            "Valor atual (€)",
            "Valorização vs custo (€)",
            "ROI vs custo (%)",
            "Valorização vs PVP (€)",
            "ROI vs PVP (%)",
            "Valor atualizado em",
            "Valor de venda (€)",
            "Data de venda",
            "Notas",
            "É Fs",
            "Potencial presente",
        ],
        rows,
        {
            5: DATE_FORMAT,
            6: DATE_FORMAT,
            19: DATE_FORMAT,
            21: MONEY_FORMAT,
            22: MONEY_FORMAT,
            23: MONEY_FORMAT,
            24: MONEY_FORMAT,
            25: PERCENT_FORMAT,
            26: MONEY_FORMAT,
            27: PERCENT_FORMAT,
            28: DATE_FORMAT,
            29: MONEY_FORMAT,
            30: DATE_FORMAT,
        },
    )


def _sets_sheet(
    db: DbSession,
    workbook: Workbook,
    scope_ids: list[uuid.UUID],
    entity_names: dict[uuid.UUID, str],
) -> None:
    models = list(
        db.scalars(
            select(LegoSetModel)
            .where(LegoSetModel.entity_id.in_(scope_ids), LegoSetModel.is_deleted.is_(False))
            .order_by(LegoSetModel.name)
        ).all()
    )
    counts = lego_service.copy_counts(db, [m.id for m in models])
    stale_days = settings_service.stale_value_days(db)
    today = dt.date.today()

    rows: list[list[Any]] = []
    for model in models:
        age = (today - model.value_updated_at).days if model.value_updated_at else None
        rows.append(
            [
                model.set_number or ("MOC" if model.is_custom else ""),
                model.name,
                model.theme,
                model.subtheme,
                model.release_date,
                model.retirement_date,
                "Sim" if model.is_retired else "Não",
                model.piece_count,
                model.minifig_count,
                entity_names.get(model.entity_id, ""),
                counts.get(model.id, 0),
                model.rrp_eur,
                model.current_value_eur,
                appreciation_eur(model.rrp_eur, model.current_value_eur),
                roi_pct(model.rrp_eur, model.current_value_eur),
                model.value_updated_at,
                "Sim" if age is None or age > stale_days else "Não",
                model.short_description or "",
                model.notes or "",
            ]
        )

    _write_sheet(
        workbook,
        "Conjuntos",
        [
            "Número",
            "Nome",
            "Tema",
            "Subtema",
            "Lançamento",
            "Retirado em",
            "Retirado",
            "Peças",
            "Minifiguras",
            "Entidade",
            "Cópias na coleção",
            "PVP original (€)",
            "Valor atual (€)",
            "Valorização vs PVP (€)",
            "ROI vs PVP (%)",
            "Valor atualizado em",
            "Valor desatualizado",
            "Descrição",
            "Notas",
        ],
        rows,
        {
            5: DATE_FORMAT,
            6: DATE_FORMAT,
            12: MONEY_FORMAT,
            13: MONEY_FORMAT,
            14: MONEY_FORMAT,
            15: PERCENT_FORMAT,
            16: DATE_FORMAT,
        },
    )


def _locations_sheet(db: DbSession, workbook: Workbook) -> None:
    locations = lego_service.list_storage_locations(db)
    rows: list[list[Any]] = [
        [
            location.area,
            location.container or "",
            location.description or "",
            location.stored_count,
            location.stored_value_eur or ZERO,
            location.capacity_pct,
            location.remaining_capacity_pct,
            "Sim" if location.is_full else "Não",
        ]
        for location in locations
    ]
    _write_sheet(
        workbook,
        "Arrumação",
        [
            "Área",
            "Contentor",
            "Descrição",
            "Cópias guardadas",
            "Valor guardado (€)",
            "Ocupação (%)",
            "Espaço livre (%)",
            "Cheio",
        ],
        rows,
        {5: MONEY_FORMAT},
    )


def _helpers_sheet(workbook: Workbook) -> None:
    """Reference lists for the enum columns the bulk importer accepts (ADR-0048)
    — the exact labels it matches, for someone filling the sheet by hand."""
    columns = [
        ("Estado de construção", list(BUILD_STATE_PT.values())),
        ("Condição", list(CONDITION_PT.values())),
        ("Origem", list(SOURCE_PT.values())),
    ]
    height = max(len(values) for _, values in columns)
    rows = [[values[i] if i < len(values) else "" for _, values in columns] for i in range(height)]
    _write_sheet(workbook, "Auxiliares", [label for label, _ in columns], rows, {})


def filename(active_entity_id: uuid.UUID | None) -> str:
    suffix = "" if active_entity_id is None else "-entidade"
    return f"colecao-lego{suffix}-{dt.date.today():%Y%m%d}.xlsx"
