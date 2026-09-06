"""Bulk import of LEGO copies and storage locations from a spreadsheet (M9.5).

Columns are matched **by header name**, not by position, against the same PT
labels `lego_export.py` writes — so the "Exportar" workbook is itself a valid
"Importar em lote" input, for either sheet. Unknown/extra export columns (name,
theme, ROI, ...) are simply ignored: everything about the *set* still comes from
Brickset, exactly as it would through the manual "Procurar" button. See
docs/decisions/0048-bulk-import-mirrors-the-export.md.

Two different consistency rules on purpose:
- Instance rows have no natural key (a copy is a copy), so importing never
  deduplicates — re-importing the same export doubles every copy, same as
  registering it twice by hand.
- Storage rows are upserted on (area, container), so re-importing the exported
  "Arrumação" sheet updates capacities/descriptions in place instead of
  duplicating locations.
"""

from __future__ import annotations

import contextlib
import datetime as dt
import io
import re
import unicodedata
import uuid
from collections.abc import Iterator
from decimal import Decimal, InvalidOperation
from typing import Any

import httpx
import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session as DbSession

from app.core.errors import AppError, NotFound, ValidationError
from app.models.household import Entity
from app.models.lego import ACQUISITION_SOURCES, BUILD_STATES, CONDITIONS, StorageLocation
from app.schemas.lego import (
    BulkImportRow,
    BulkImportRowResult,
    LegoSetInstanceCreate,
    LegoSetModelCreate,
    LegoSetModelUpdate,
    StorageBulkImportError,
    StorageBulkImportOut,
    StorageLocationCreate,
    StorageLocationUpdate,
)
from app.services import lego_brickset_jobs, lego_provider, lego_service
from app.services.lego_export import BUILD_STATE_PT, CONDITION_PT, SOURCE_PT

INSTANCE_SHEET_NAMES = ("Cópias", "Copias")
STORAGE_SHEET_NAMES = ("Arrumação", "Arrumacao")

INSTANCE_HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "set_number": ("numero", "lego id", "id do conjunto"),
    "entity": ("entidade",),
    "area": ("area",),
    "container": ("contentor",),
    "build_state": ("estado de construcao",),
    "condition": ("condicao",),
    "has_box": ("tem caixa",),
    "has_instructions": ("tem instrucoes", "tem manual"),
    "is_fs": ("e fs",),
    "missing_parts": ("pecas em falta",),
    "acquisition_date": ("data de aquisicao",),
    "acquisition_source": ("origem",),
    "acquisition_cost_eur": ("custo", "custo de aquisicao"),
    "current_value_eur": ("preco atual", "valor atual"),
    "notes": ("notas",),
}

STORAGE_HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "area": ("area",),
    "container": ("contentor",),
    "description": ("descricao",),
    "capacity_pct": ("ocupacao",),
}


def _norm(text: str) -> str:
    ascii_text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", ascii_text.lower()).strip()


_REVERSE_SOURCE_PT = {_norm(label): code for code, label in SOURCE_PT.items()}
_REVERSE_BUILD_STATE_PT = {_norm(label): code for code, label in BUILD_STATE_PT.items()}
_REVERSE_CONDITION_PT = {_norm(label): code for code, label in CONDITION_PT.items()}


# --- Sheet reading -------------------------------------------------------------
def _load_rows(
    data: bytes, *, sheet_names: tuple[str, ...], aliases: dict[str, tuple[str, ...]]
) -> list[dict[str, Any]]:
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:
        raise ValidationError("Não foi possível ler o ficheiro Excel.") from exc

    sheet = None
    for wanted in sheet_names:
        match = next((name for name in workbook.sheetnames if _norm(name) == _norm(wanted)), None)
        if match is not None:
            sheet = workbook[match]
            break
    if sheet is None:
        sheet = workbook[workbook.sheetnames[0]]

    iterator = sheet.iter_rows(values_only=True)
    header_row = next(iterator, None)
    if header_row is None:
        raise ValidationError("A folha está vazia.")

    columns: dict[str, int] = {}
    for index, cell in enumerate(header_row):
        if cell is None:
            continue
        normalized = _norm(str(cell))
        for field_name, alias_list in aliases.items():
            if field_name not in columns and normalized in alias_list:
                columns[field_name] = index

    rows: list[dict[str, Any]] = []
    for row_number, values in enumerate(iterator, start=2):
        if all(value in (None, "") for value in values):
            continue
        row: dict[str, Any] = {"_row": row_number}
        for field_name, index in columns.items():
            row[field_name] = values[index] if index < len(values) else None
        rows.append(row)
    return rows


# --- Cell parsing ---------------------------------------------------------------
def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _decimal(value: Any) -> Decimal | None:
    if value in (None, ""):
        return None
    try:
        return Decimal(str(value).strip().replace(",", ".")).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        return None


def _date(value: Any) -> tuple[dt.date | None, str | None]:
    if value in (None, ""):
        return None, None
    if isinstance(value, dt.datetime):
        return value.date(), None
    if isinstance(value, dt.date):
        return value, None
    text = str(value).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return dt.datetime.strptime(text, fmt).date(), None
        except ValueError:
            continue
    return None, f"Data «{text}» não reconhecida."


def _boolean(value: Any, *, default: bool) -> bool:
    text = _text(value)
    if text is None:
        return default
    normalized = _norm(text)
    if normalized in ("sim", "true", "1", "yes"):
        return True
    if normalized in ("nao", "false", "0", "no"):
        return False
    return default


def _label(
    value: Any, reverse: dict[str, str], codes: tuple[str, ...], field_label: str
) -> tuple[str | None, str | None]:
    text = _text(value)
    if text is None:
        return None, None
    if text.upper() in codes:
        return text.upper(), None
    code = reverse.get(_norm(text))
    if code is None:
        return None, f"{field_label} «{text}» não reconhecida."
    return code, None


def _integer_pct(value: Any) -> int | None:
    if value in (None, ""):
        return None
    try:
        number = round(float(str(value).strip().replace(",", ".")))
    except ValueError:
        return None
    return max(0, min(100, number))


# --- Instances: preview ---------------------------------------------------------
def preview_instances(
    db: DbSession, *, data: bytes, household_id: uuid.UUID, active_entity_id: uuid.UUID | None
) -> list[BulkImportRow]:
    raw_rows = _load_rows(data, sheet_names=INSTANCE_SHEET_NAMES, aliases=INSTANCE_HEADER_ALIASES)
    if not raw_rows:
        raise ValidationError("Não há linhas para importar.")

    entities = {
        _norm(name): (entity_id, is_readonly)
        for entity_id, name, is_readonly in db.execute(
            select(Entity.id, Entity.name, Entity.is_readonly).where(
                Entity.household_id == household_id, Entity.is_deleted.is_(False)
            )
        ).all()
    }
    locations = {
        (_norm(area), _norm(container or "")): location_id
        for location_id, area, container in db.execute(
            select(StorageLocation.id, StorageLocation.area, StorageLocation.container).where(
                StorageLocation.is_deleted.is_(False)
            )
        ).all()
    }

    return [
        _resolve_instance_row(
            raw, entities=entities, locations=locations, active_entity_id=active_entity_id
        )
        for raw in raw_rows
    ]


def _resolve_instance_row(
    raw: dict[str, Any],
    *,
    entities: dict[str, tuple[uuid.UUID, bool]],
    locations: dict[tuple[str, str], uuid.UUID],
    active_entity_id: uuid.UUID | None,
) -> BulkImportRow:
    errors: dict[str, str] = {}

    set_number = _text(raw.get("set_number"))
    if not set_number:
        errors["set_number"] = "Indique o número do conjunto."

    entity_name = _text(raw.get("entity"))
    entity_id: uuid.UUID | None = None
    if entity_name:
        match = entities.get(_norm(entity_name))
        if match is None:
            errors["entity_id"] = f"Entidade «{entity_name}» não encontrada."
        elif match[1]:
            errors["entity_id"] = f"«{entity_name}» está em modo de leitura."
        else:
            entity_id = match[0]
    elif active_entity_id is not None:
        entity_id = active_entity_id
    else:
        errors["entity_id"] = "Indique a entidade."

    area = _text(raw.get("area"))
    container = _text(raw.get("container"))
    storage_location_id: uuid.UUID | None = None
    if area:
        found = locations.get((_norm(area), _norm(container or "")))
        if found is None:
            label = f"{area} › {container}" if container else area
            errors["storage_location_id"] = f"Local «{label}» não encontrado."
        else:
            storage_location_id = found
    elif container:
        errors["storage_location_id"] = "Indique também a área."

    raw_cost = raw.get("acquisition_cost_eur")
    cost = _decimal(raw_cost)
    if raw_cost not in (None, "") and cost is None:
        errors["acquisition_cost_eur"] = "Custo inválido."

    raw_current_value = raw.get("current_value_eur")
    current_value = _decimal(raw_current_value)
    if raw_current_value not in (None, "") and current_value is None:
        errors["current_value_eur"] = "Preço atual inválido."

    date_value, date_error = _date(raw.get("acquisition_date"))
    if date_error:
        errors["acquisition_date"] = date_error

    source, source_error = _label(
        raw.get("acquisition_source"), _REVERSE_SOURCE_PT, ACQUISITION_SOURCES, "Origem"
    )
    if source_error:
        errors["acquisition_source"] = source_error

    build_state, build_state_error = _label(
        raw.get("build_state"), _REVERSE_BUILD_STATE_PT, BUILD_STATES, "Estado de construção"
    )
    if build_state_error:
        errors["build_state"] = build_state_error

    condition, condition_error = _label(
        raw.get("condition"), _REVERSE_CONDITION_PT, CONDITIONS, "Condição"
    )
    if condition_error:
        errors["condition"] = condition_error

    return BulkImportRow(
        row_number=raw["_row"],
        set_number=set_number.upper() if set_number else None,
        entity_id=entity_id,
        entity_name=entity_name,
        storage_location_id=storage_location_id,
        storage_area=area,
        storage_container=container,
        acquisition_cost_eur=cost or Decimal("0.00"),
        acquisition_date=date_value,
        acquisition_source=source,  # type: ignore[arg-type]
        build_state=build_state,  # type: ignore[arg-type]
        condition=condition,  # type: ignore[arg-type]
        has_box=_boolean(raw.get("has_box"), default=True),
        has_instructions=_boolean(raw.get("has_instructions"), default=True),
        is_fs=_boolean(raw.get("is_fs"), default=False),
        missing_parts=_text(raw.get("missing_parts")),
        notes=_text(raw.get("notes")),
        current_value_eur=current_value,
        errors=errors,
    )


# --- Instances: commit -----------------------------------------------------------
def commit_instances(
    db: DbSession, rows: list[BulkImportRow], *, household_id: uuid.UUID, actor_user_id: uuid.UUID
) -> list[BulkImportRowResult]:
    return list(
        iter_commit_instances(db, rows, household_id=household_id, actor_user_id=actor_user_id)
    )


def iter_commit_instances(
    db: DbSession, rows: list[BulkImportRow], *, household_id: uuid.UUID, actor_user_id: uuid.UUID
) -> Iterator[BulkImportRowResult]:
    """Same row-by-row commit as `commit_instances`, yielded as each row finishes —
    lets the router stream progress instead of making the client wait for the
    whole batch (a Brickset lookup per new set can make this a slow request).

    Commits (or rolls back) right here, per row — the router returns a
    `StreamingResponse` immediately, so the request-scoped session's own
    commit-on-teardown fires before this generator is ever iterated and would
    otherwise never persist anything written while streaming (losing exactly
    the last row committed here, and any row before one that never reached a
    `db.commit()` of its own — e.g. `queue_brickset_assets`)."""
    for row in rows:
        try:
            result = _commit_instance_row(
                db, row, household_id=household_id, actor_user_id=actor_user_id
            )
            db.commit()
            yield result
        except AppError as exc:
            db.rollback()
            yield BulkImportRowResult(row_number=row.row_number, ok=False, message=str(exc.detail))
        except Exception as exc:  # one bad row must not sink the whole batch
            db.rollback()
            yield BulkImportRowResult(row_number=row.row_number, ok=False, message=str(exc))


def _commit_instance_row(
    db: DbSession, row: BulkImportRow, *, household_id: uuid.UUID, actor_user_id: uuid.UUID
) -> BulkImportRowResult:
    if row.errors:
        return BulkImportRowResult(
            row_number=row.row_number, ok=False, message="Linha por corrigir."
        )
    if not row.set_number:
        return BulkImportRowResult(
            row_number=row.row_number, ok=False, message="Sem número de conjunto."
        )
    if row.entity_id is None:
        return BulkImportRowResult(row_number=row.row_number, ok=False, message="Sem entidade.")

    entity = db.get(Entity, row.entity_id)
    if entity is None or entity.is_deleted or entity.household_id != household_id:
        return BulkImportRowResult(
            row_number=row.row_number, ok=False, message="Entidade desconhecida."
        )
    if entity.is_readonly:
        return BulkImportRowResult(
            row_number=row.row_number, ok=False, message="Entidade em modo de leitura."
        )

    if row.storage_location_id is not None:
        try:
            lego_service.get_storage_location(db, row.storage_location_id)
        except NotFound:
            return BulkImportRowResult(
                row_number=row.row_number, ok=False, message="Local de arrumação desconhecido."
            )

    set_number = row.set_number.strip().upper()
    model = lego_service.find_model_by_set_number(db, entity.id, set_number)
    if model is None:
        provider = lego_provider.get_provider(db)
        if not getattr(provider, "enabled", False):
            return BulkImportRowResult(
                row_number=row.row_number,
                ok=False,
                message=(
                    f"{set_number} é novo e o Brickset está desligado — "
                    "registe-o manualmente primeiro."
                ),
            )
        try:
            lookup = provider.lookup(set_number)
        except (httpx.HTTPError, ValueError) as exc:
            return BulkImportRowResult(
                row_number=row.row_number,
                ok=False,
                message=f"Brickset indisponível ({exc.__class__.__name__}).",
            )
        if not lookup.found:
            return BulkImportRowResult(
                row_number=row.row_number,
                ok=False,
                message=lookup.message or f"{set_number} não encontrado no Brickset.",
            )
        draft = LegoSetModelCreate(
            set_number=set_number,
            is_custom=False,
            name=lookup.name or set_number,
            theme=lookup.theme,
            subtheme=lookup.subtheme,
            release_date=lookup.release_date,
            retirement_date=lookup.retirement_date,
            piece_count=lookup.piece_count,
            minifig_count=lookup.minifig_count,
            age_min=lookup.age_min,
            age_max=lookup.age_max,
            box_height_cm=lookup.box_height_cm,
            box_width_cm=lookup.box_width_cm,
            box_depth_cm=lookup.box_depth_cm,
            box_weight_kg=lookup.box_weight_kg,
            rrp_eur=lookup.rrp_eur,
            short_description=lookup.short_description,
            image_url=lookup.image_url,
            current_value_eur=row.current_value_eur,
        )
        model = lego_service.create_model(
            db, draft, entity_id=entity.id, actor_user_id=actor_user_id
        )
        # The gallery/manuals are a bonus, never worth failing the row over —
        # and slow enough (one HTTP download each) that they run as background
        # jobs instead of blocking this row's result (ADR-0049).
        with contextlib.suppress(ValidationError):
            lego_brickset_jobs.queue_brickset_assets(db, model, actor_user_id=actor_user_id)
    elif row.current_value_eur is not None and (
        model.current_value_eur is None or row.current_value_eur > model.current_value_eur
    ):
        # Copies of the same set can carry different "preço atual" values across
        # rows (e.g. a hand-edited sheet) — the highest one wins, and it only
        # ever pushes the value up, never down, across the whole batch.
        lego_service.update_model(
            db,
            model,
            LegoSetModelUpdate(current_value_eur=row.current_value_eur),
            actor_user_id=actor_user_id,
        )

    instance_payload = LegoSetInstanceCreate(
        entity_id=entity.id,
        lego_set_model_id=model.id,
        acquisition_date=row.acquisition_date,
        acquisition_cost_eur=row.acquisition_cost_eur,
        acquisition_source=row.acquisition_source,
        storage_location_id=row.storage_location_id,
        build_state=row.build_state,
        condition=row.condition,
        has_box=row.has_box,
        has_instructions=row.has_instructions,
        is_fs=row.is_fs,
        missing_parts=row.missing_parts,
        notes=row.notes,
    )
    instance = lego_service.create_instance(
        db, instance_payload, entity_id=entity.id, actor_user_id=actor_user_id
    )
    return BulkImportRowResult(
        row_number=row.row_number,
        ok=True,
        message="Cópia registada.",
        lego_set_instance_id=instance.id,
    )


# --- Storage locations: one-step upsert -----------------------------------------
def import_storage_locations(
    db: DbSession, *, data: bytes, actor_user_id: uuid.UUID
) -> StorageBulkImportOut:
    raw_rows = _load_rows(data, sheet_names=STORAGE_SHEET_NAMES, aliases=STORAGE_HEADER_ALIASES)
    created = 0
    updated = 0
    errors: list[StorageBulkImportError] = []

    for raw in raw_rows:
        area = _text(raw.get("area"))
        if not area:
            errors.append(StorageBulkImportError(row_number=raw["_row"], message="Indique a área."))
            continue
        container = _text(raw.get("container"))
        description = _text(raw.get("description"))
        capacity_pct = _integer_pct(raw.get("capacity_pct"))

        existing = db.scalar(
            select(StorageLocation).where(
                StorageLocation.area == area,
                StorageLocation.container.is_not_distinct_from(container),
                StorageLocation.is_deleted.is_(False),
            )
        )
        try:
            if existing is not None:
                lego_service.update_storage_location(
                    db,
                    existing,
                    StorageLocationUpdate(
                        description=(
                            description if description is not None else existing.description
                        ),
                        capacity_pct=(
                            capacity_pct if capacity_pct is not None else existing.capacity_pct
                        ),
                    ),
                    actor_user_id=actor_user_id,
                )
                updated += 1
            else:
                lego_service.create_storage_location(
                    db,
                    StorageLocationCreate(
                        area=area,
                        container=container,
                        description=description,
                        capacity_pct=capacity_pct,
                    ),
                    actor_user_id=actor_user_id,
                )
                created += 1
        except AppError as exc:
            errors.append(StorageBulkImportError(row_number=raw["_row"], message=str(exc.detail)))

    return StorageBulkImportOut(created=created, updated=updated, errors=errors)
